from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


def export_word_from_markdown(markdown_content: str, file_path: str, title: str | None = None) -> str:
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")

    doc = Document()
    _set_default_fonts(doc)

    if title:
        doc.add_heading(title, level=0)

    lines = markdown_content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if _looks_like_table_start(lines, index):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            _add_markdown_table(doc, table_lines)
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith(("- ", "* ")):
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_markdown_runs(paragraph, stripped[2:].strip())
        elif re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            _add_markdown_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
        else:
            paragraph = doc.add_paragraph()
            _add_markdown_runs(paragraph, stripped)

        index += 1

    doc.save(path)
    return str(path)


def _set_default_fonts(doc: Document) -> None:
    styles = doc.styles
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        rpr = style._element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")


def _add_markdown_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        run.bold = bold
        run.font.name = "Times New Roman"
        rpr = run._element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")


def _looks_like_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    next_line = lines[index + 1].strip()
    return current.startswith("|") and next_line.startswith("|") and "---" in next_line


def _add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    if len(table_lines) < 2:
        return

    rows = [_split_table_row(line) for line in table_lines if not _is_table_separator(line)]
    if not rows:
        return

    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for col_index, value in enumerate(rows[0]):
        table.rows[0].cells[col_index].text = value

    for row in rows[1:]:
        cells = table.add_row().cells
        for col_index, value in enumerate(row[: len(cells)]):
            cells[col_index].text = value


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    stripped = line.strip().strip("|").replace("|", "").replace(":", "").replace("-", "").strip()
    return stripped == ""
